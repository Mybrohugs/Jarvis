"""Excel operator on temporary stand-ins for Documents/Desktop/Downloads (never the real folders)."""
import datetime
from pathlib import Path
import tempfile
import unittest
import openpyxl
from openpyxl.chart import BarChart, Reference
from developer_core.jarvis_local.excel_ops import ExcelOps
from developer_core.jarvis_local.explorer import Explorer
from developer_core.jarvis_local.intent import understand
from developer_core.jarvis_local.phone_executor import validate
from developer_core.jarvis_local.plan_runner import run_plan
from developer_core.jarvis_local.router import UnifiedTaskRouter
from developer_core.jarvis_local.session import Denied, TrustedLocalSession


class ExcelTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        base = Path(tmp.name)
        self.docs, self.desk, self.dl, ws = base / 'Documents', base / 'Desktop', base / 'Downloads', base / 'ws'
        for d in (self.docs, self.desk, self.dl, ws):
            d.mkdir()
        self.asked, self.opened = [], []

        def confirmer(summary, challenge):
            self.asked.append(summary)
            return True
        self.router = UnifiedTaskRouter(TrustedLocalSession.for_tests(confirmer), ws)
        self.router._explorer = Explorer(self.router.ws, roots={'Documents': self.docs, 'Desktop': self.desk,
                                                               'Downloads': self.dl})
        self.router._excel_ops = ExcelOps(self.router.ws, explorer=self.router._explorer,
                                          start=lambda p: self.opened.append(p),
                                          titles=lambda: ['Расходы.xlsx - Excel'], sleep=lambda s: None)

    def run_phrase(self, phrase):
        return run_plan(self.router, understand(phrase))

    def rows(self, path):
        wb = openpyxl.load_workbook(path)
        try:
            return [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()

    def make(self, folder, name, rows, chart=False):
        wb = openpyxl.Workbook()
        for r in rows:
            wb.active.append(r)
        if chart:
            c = BarChart()
            c.add_data(Reference(wb.active, min_col=2, min_row=1, max_row=len(rows)), titles_from_data=True)
            wb.active.add_chart(c, 'E2')
        p = folder / name
        wb.save(p)
        return p

    def test_create_add_open_in_one_phrase(self):
        out = self.run_phrase('создай таблицу Расходы с колонками Дата, Сумма, Категория на рабочем столе '
                              'и добавь строку «29.10.2026; 500; еда» и открой её')
        self.assertEqual(out['status'], 'PASS', out)
        rows = self.rows(self.desk / 'Расходы.xlsx')
        self.assertEqual(rows[0], ['Дата', 'Сумма', 'Категория'])
        self.assertEqual(rows[1], [datetime.datetime(2026, 10, 29), 500, 'еда'])
        self.assertEqual(self.opened, [str(self.desk / 'Расходы.xlsx')])
        self.assertTrue(out['completed'][2]['result']['verified'])

    def test_column_math_and_find(self):
        self.make(self.docs, 'Расходы.xlsx', [['Дата', 'Сумма', 'Категория'], ['01.10', 500, 'еда'],
                                              ['02.10', '1 200,50', 'такси'], ['03.10', 300, 'еда']])
        s = self.run_phrase('посчитай сумму по столбцу Сумма в таблице расходы')['completed'][0]['result']
        self.assertAlmostEqual(s['value'], 2000.5)
        self.assertIn('2000,50', s['speech'])
        avg = self.router.dispatch('excel-column', name='расходы', column='сумм', func='max')
        self.assertEqual(avg['value'], 1200.5)
        found = self.run_phrase('найди такси в таблице расходы')['completed'][0]['result']
        self.assertEqual([r['row'] for r in found['rows']], [3])
        summary = self.run_phrase('что в таблице расходы')['completed'][0]['result']
        self.assertEqual(summary['rows'], 3)
        self.assertIn('Итого: Сумма 2000,50', summary['speech'])

    def test_append_keeps_backup_and_rolls_back(self):
        p = self.make(self.docs, 'Бюджет.xlsx', [['Статья', 'Сумма'], ['Аренда', 30000]])
        before = p.read_bytes()
        out = self.run_phrase('добавь строку «Еда; 12000» в таблицу бюджет')
        self.assertEqual(out['status'], 'PASS', out)
        self.assertEqual(self.rows(p)[-1], ['Еда', 12000])
        self.assertEqual(self.asked[-1]['writes'], 'original (backup kept)')
        self.assertEqual(self.run_phrase('отмени последнее действие')['status'], 'PASS')
        self.assertEqual(p.read_bytes(), before)

    def test_rollback_refused_after_owner_edit(self):
        p = self.make(self.docs, 'Бюджет.xlsx', [['Статья', 'Сумма']])
        out = self.router.dispatch('excel-append', name='бюджет', cells=['Еда', '1'])
        self.make(self.docs, 'Бюджет.xlsx', [['Статья', 'Сумма'], ['изменено вручную', 2]])
        with self.assertRaises(Denied):
            self.router.dispatch('rollback', entry_id=out['audit_id'])
        self.assertEqual(self.rows(p)[-1], ['изменено вручную', 2])

    def test_workbook_with_chart_gets_a_copy(self):
        p = self.make(self.docs, 'Отчёт.xlsx', [['Месяц', 'Сумма'], ['Янв', 1], ['Фев', 2]], chart=True)
        before = p.read_bytes()
        out = self.router.dispatch('excel-append', name='отчёт', cells=['Мар', '3'])
        self.assertTrue(out['copy'])
        self.assertEqual(p.read_bytes(), before)  # the original with the chart is untouched
        copy = self.docs / 'Отчёт (JARVIS).xlsx'
        self.assertEqual(self.rows(copy)[-1], ['Мар', 3])
        self.router.dispatch('rollback', entry_id=out['audit_id'])
        self.assertFalse(copy.exists())

    def test_refusals(self):
        p = self.make(self.docs, 'Бюджет.xlsx', [['Статья', 'Сумма']])
        (self.docs / '~$Бюджет.xlsx').write_bytes(b'lock')  # open in Excel
        with self.assertRaises(Denied):
            self.router.dispatch('excel-append', name='бюджет', cells=['a', '1'])
        (self.docs / '~$Бюджет.xlsx').unlink()
        with self.assertRaises(Denied):  # more values than columns
            self.router.dispatch('excel-append', name='бюджет', cells=['a', '1', 'x'])
        self.make(self.desk, 'Бюджет.xlsx', [['x']])
        with self.assertRaises(Denied):  # two tables with the same name: ask which
            self.router.dispatch('excel-summary', name='бюджет')
        (self.desk / 'Бюджет.xlsx').unlink()
        with self.assertRaises(Denied):
            self.router.dispatch('excel-create', name='Бюджет', headers=['a'], folder='Documents')  # never overwrites
        with self.assertRaises(Denied):
            self.router.dispatch('excel-summary', name='нет такой')
        self.assertEqual(self.rows(p), [['Статья', 'Сумма']])

    def test_formula_from_a_phrase_stays_text(self):
        p = self.make(self.docs, 'Т.xlsx', [['A']])
        self.router.dispatch('excel-append', name='т', cells=['=HYPERLINK("http://evil")'])
        wb = openpyxl.load_workbook(p)
        cell = wb.active['A2']
        self.assertEqual(cell.data_type, 's')
        wb.close()

    def test_phone_reads_but_does_not_write(self):
        self.assertEqual(validate({'task': 'excel-column', 'args': {'name': 'расходы', 'column': 'Сумма', 'func': 'sum'}})[0],
                         'excel-column')
        for bad in ({'task': 'excel-append', 'args': {'name': 'x', 'cells': ['1']}},
                    {'task': 'excel-create', 'args': {'name': 'x', 'headers': ['a'], 'folder': 'Documents'}},
                    {'task': 'excel-column', 'args': {'name': 'x', 'column': 'y', 'func': 'drop'}}):
            with self.assertRaises(Denied):
                validate(bad)

    def test_workspace_xlsx_phrases_unchanged(self):
        self.assertEqual(understand('добавь строку «1;2» в таблицу data.xlsx')['steps'][0]['task'], 'office-plan')


if __name__ == '__main__':
    unittest.main()
