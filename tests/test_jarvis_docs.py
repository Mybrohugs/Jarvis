"""Word/PDF operator on temporary stand-ins for Documents/Desktop/Downloads (never the real folders).
PDFs are parsed by the real isolated worker; Word and Docker are replaced by fakes here."""
from pathlib import Path
import subprocess
import tempfile
import unittest
import docx
from developer_core.jarvis_local.docs_ops import DocsOps
from developer_core.jarvis_local.explorer import Explorer
from developer_core.jarvis_local.intent import understand
from developer_core.jarvis_local.phone_executor import validate
from developer_core.jarvis_local.plan_runner import run_plan
from developer_core.jarvis_local.router import UnifiedTaskRouter
from developer_core.jarvis_local.session import Denied, TrustedLocalSession


def pdf_bytes(pages):
    """A tiny valid PDF with one line of ASCII text per page."""
    objs = ['<< /Type /Catalog /Pages 2 0 R >>', None, '<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>']
    kids = []
    for text in pages:
        stream = 'BT /F1 18 Tf 72 720 Td (%s) Tj ET' % text
        objs.append('<< /Length %d >>\nstream\n%s\nendstream' % (len(stream), stream))
        objs.append('<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents %d 0 R '
                    '/Resources << /Font << /F1 3 0 R >> >> >>' % (len(objs)))
        kids.append('%d 0 R' % len(objs))
    objs[1] = '<< /Type /Pages /Kids [%s] /Count %d >>' % (' '.join(kids), len(kids))
    out, offsets = b'%PDF-1.4\n', []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += ('%d 0 obj\n%s\nendobj\n' % (i, o)).encode('latin-1')
    xref = len(out)
    out += ('xref\n0 %d\n0000000000 65535 f \n' % (len(objs) + 1)).encode()
    out += b''.join(('%010d 00000 n \n' % o).encode() for o in offsets)
    out += ('trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n' % (len(objs) + 1, xref)).encode()
    return out


class DocsTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        base = Path(tmp.name)
        self.docs, self.desk, self.dl, ws = base / 'Documents', base / 'Desktop', base / 'Downloads', base / 'ws'
        for d in (self.docs, self.desk, self.dl, ws):
            d.mkdir()
        self.asked, self.opened, self.sent = [], [], []

        def confirmer(summary, challenge):
            self.asked.append(summary)
            return True
        self.router = UnifiedTaskRouter(TrustedLocalSession.for_tests(confirmer), ws)
        self.router._explorer = Explorer(self.router.ws, roots={'Documents': self.docs, 'Desktop': self.desk,
                                                               'Downloads': self.dl})
        real_run = subprocess.run

        def run(cmd, **kw):  # Word and Docker are faked; the PDF worker runs for real
            if 'topdf' in cmd:
                Path(cmd[-1]).write_bytes(pdf_bytes(['Converted by Word']))
                return real_run(cmd[:-3] + ['info', cmd[-1]], **kw)
            if cmd[1:2] == ['run']:
                work = Path(cmd[cmd.index('-v') + 1].rsplit(':/work', 1)[0])
                (work / 'out.pdf').write_bytes(pdf_bytes(['Scanned passport text now searchable']))
                return subprocess.CompletedProcess(cmd, 0, b'', b'')
            return real_run(cmd, **kw)
        self.router._docs_ops = DocsOps(self.router.ws, explorer=self.router._explorer, run=run,
                                        start=lambda p: self.opened.append(p), titles=lambda: ['Письмо.docx - Word'],
                                        sleep=lambda s: None,
                                        ask=lambda system, user: self.sent.append(user) or 'Это письмо о сроках отчёта.')

    def run_phrase(self, phrase):
        return run_plan(self.router, understand(phrase))

    def test_create_then_pdf_then_open_in_one_phrase(self):
        out = self.run_phrase('создай документ «Письмо» с текстом «Добрый день! Высылаю отчёт.» на рабочем столе '
                              'и сделай из него pdf')
        self.assertEqual(out['status'], 'PASS', out)
        paras = [p.text for p in docx.Document(str(self.desk / 'Письмо.docx')).paragraphs]
        self.assertEqual(paras, ['Письмо', 'Добрый день! Высылаю отчёт.'])
        self.assertTrue((self.desk / 'Письмо.pdf').read_bytes().startswith(b'%PDF-'))
        self.assertEqual(out['completed'][1]['result']['pages'], 1)
        self.assertEqual(self.run_phrase('открой документ Письмо')['completed'][0]['result']['verified'], True)
        undo = self.router.dispatch('rollback', entry_id=out['completed'][1]['result']['audit_id'])
        self.assertEqual(undo['status'], 'PASS')
        self.assertFalse((self.desk / 'Письмо.pdf').exists())  # quarantined, not deleted
        self.assertEqual(len(list((self.router.ws.state / 'quarantine').rglob('Письмо.pdf'))), 1)

    def test_read_find_and_retell(self):
        (self.docs / 'Отчёт за май.pdf').write_bytes(pdf_bytes(['Invoice total 500 rub', 'Penalty clause on page two']))
        r = self.run_phrase('сколько страниц в pdf отчёт за май')['completed'][0]['result']
        self.assertEqual(r['pages'], 2)
        self.assertIn('2 страницы', r['speech'])
        f = self.run_phrase('найди «penalty» в документе отчёт за май')['completed'][0]['result']
        self.assertEqual(f['count'], 1)
        d = docx.Document()
        d.add_paragraph('Срок сдачи отчёта перенесён на пятницу, ответственный Иванов.')
        d.save(str(self.docs / 'Письмо.docx'))
        s = self.run_phrase('перескажи документ письмо')['completed'][0]['result']
        self.assertEqual(s['speech'], 'Это письмо о сроках отчёта.')
        self.assertIn('ответственный Иванов', self.sent[-1])
        self.assertEqual(self.asked[-1]['task'], 'doc-summary')  # the text leaves the PC only after approval

    def test_append_backup_and_rollback(self):
        d = docx.Document()
        d.add_paragraph('Добрый день!')
        p = self.docs / 'Письмо.docx'
        d.save(str(p))
        before = p.read_bytes()
        out = self.run_phrase('допиши в документ Письмо «С уважением, Роман»')
        self.assertEqual(out['status'], 'PASS', out)
        self.assertEqual(docx.Document(str(p)).paragraphs[-1].text, 'С уважением, Роман')
        self.assertEqual(self.run_phrase('отмени последнее действие')['status'], 'PASS')
        self.assertEqual(p.read_bytes(), before)

    def test_append_refusals(self):
        d = docx.Document()
        d.add_paragraph('x')
        d.save(str(self.docs / 'Договор аренды.docx'))
        (self.docs / '~$говор аренды.docx').write_bytes(b'lock')  # Word cuts the first two letters
        with self.assertRaises(Denied):
            self.router.dispatch('doc-append', name='договор аренды', text='y')
        with self.assertRaises(Denied):
            self.router.dispatch('doc-create', name='Договор аренды', text='', folder='Documents')  # never overwrites

    def test_merge_and_ocr(self):
        (self.docs / 'Акт.pdf').write_bytes(pdf_bytes(['Act']))
        (self.docs / 'Счёт.pdf').write_bytes(pdf_bytes(['Bill', 'Bill page 2']))
        out = self.run_phrase('объедини pdf «Акт» и «Счёт»')
        self.assertEqual(out['status'], 'PASS', out)
        self.assertEqual(out['completed'][0]['result']['pages'], 3)
        self.assertTrue((self.docs / 'Акт + Счёт.pdf').is_file())
        (self.docs / 'Скан паспорта.pdf').write_bytes(pdf_bytes(['']))
        r = self.run_phrase('прочитай pdf скан паспорта')['completed'][0]['result']
        self.assertTrue(r['scan'])
        self.router._docs_ops.ocr.__func__  # noqa: B018 (the method exists)
        import os
        os.environ.setdefault('LOCALAPPDATA', str(self.docs.parent))
        docker = Path(os.environ['LOCALAPPDATA']) / 'Programs' / 'DockerDesktop' / 'resources' / 'bin' / 'docker.exe'
        if docker.is_file():
            o = self.run_phrase('распознай pdf скан паспорта')
            self.assertEqual(o['status'], 'PASS', o)
            self.assertTrue((self.docs / 'Скан паспорта (текст).pdf').is_file())

    def test_broken_pdf_does_not_hurt(self):
        (self.docs / 'Битый.pdf').write_bytes(b'%PDF-1.4 garbage')
        with self.assertRaises(Denied):
            self.router.dispatch('doc-read', name='битый')

    def test_phone_reads_only(self):
        self.assertEqual(validate({'task': 'doc-find', 'args': {'name': 'договор', 'text': 'штраф'}})[0], 'doc-find')
        for bad in ({'task': 'doc-append', 'args': {'name': 'x', 'text': 'y'}},
                    {'task': 'doc-summary', 'args': {'name': 'x'}}, {'task': 'pdf-merge', 'args': {'names': ['a', 'b']}}):
            with self.assertRaises(Denied):
                validate(bad)


if __name__ == '__main__':
    unittest.main()
