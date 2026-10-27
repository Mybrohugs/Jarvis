"""Speech + presentation (no network): Claude is replaced by the sample plan."""
import io
import json
from pathlib import Path
import tempfile
import unittest
from pptx import Presentation
from developer_core.jarvis_local.claude_writer import KINDS, SCHEMA, SLIDE, brief, validate_plan
from developer_core.jarvis_local.router import UnifiedTaskRouter
from developer_core.jarvis_local.session import Denied, TrustedLocalSession

SAMPLE = Path(__file__).parent / 'data' / 'speech_plan_sample.json'


def sample_planner(topic, minutes, **kw):
    return validate_plan(json.loads(SAMPLE.read_text(encoding='utf-8')))


class FakePictures:
    name = 'fake'

    def __init__(self, raw=None):
        if raw is None:
            from PIL import Image
            buf = io.BytesIO()
            Image.new('RGB', (1200, 800), (30, 90, 160)).save(buf, 'PNG')
            raw = buf.getvalue()
        self.raw = raw

    def fetch(self, description):
        return self.raw, 'Фото: тест'


class SpeechTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.summaries = []

    def router(self, approve=True, read_only=False, images=()):
        session = TrustedLocalSession.read_only() if read_only else TrustedLocalSession.for_tests(
            lambda summary, challenge: self.summaries.append(summary) or approve)
        r = UnifiedTaskRouter(session, self.tmp.name)
        r.speech_planner = sample_planner
        r.speech_images = list(images)
        return r

    def test_pictures_go_only_to_marked_slides(self):
        out = self.router(images=[FakePictures()]).dispatch('speech-create', topic='ИИ', minutes=5, themes=('hud',))
        self.assertEqual([i['slide'] for i in out['images'] if i['ok']], [1, 5, 8])
        self.assertEqual(self.summaries[0]['pictures_from'], ['fake'])
        deck = Presentation(str(Path(out['folder']) / 'slides-hud.pptx'))
        with_pictures = [i for i, s in enumerate(deck.slides, 1) if any(sh.shape_type == 13 for sh in s.shapes)]
        self.assertEqual(with_pictures, [1, 5, 8])
        self.assertIn('[Фото: тест]', deck.slides[0].notes_slide.notes_text_frame.text)
        self.assertEqual(len([f for f in out['files'] if f.endswith('.jpg')]), 3)  # rollback covers pictures

    def test_bad_download_is_skipped_not_fatal(self):
        bad = FakePictures(raw=b'<html>not an image</html>')
        out = self.router(images=[bad]).dispatch('speech-create', topic='ИИ', minutes=5, themes=('hud',))
        self.assertEqual(out['status'], 'PASS')
        self.assertFalse(any(i['ok'] for i in out['images']))
        self.assertIn('Not a valid image', out['images'][0]['error'])

    def test_creates_speech_and_three_decks_with_notes(self):
        r = self.router()
        out = r.dispatch('speech-create', topic='ИИ в логистике', minutes=10, themes=('hud', 'business', 'creative'))
        self.assertEqual(out['status'], 'PASS')
        self.assertIn('Anthropic', self.summaries[0]['sent_to'])  # the user sees where the topic goes
        folder = Path(out['folder'])
        names = sorted(p.name for p in folder.iterdir())
        self.assertEqual(names, ['plan.json', 'slides-business.pptx', 'slides-creative.pptx', 'slides-hud.pptx',
                                 'speech.docx'])
        deck = Presentation(str(folder / 'slides-hud.pptx'))
        self.assertEqual(len(deck.slides), 8)
        self.assertIn('вторник', deck.slides[1].notes_slide.notes_text_frame.text)
        self.assertEqual(out['to_verify'], ['Доля времени диспетчера на ручное планирование'])

    def test_rollback_quarantines_all_created_files(self):
        r = self.router()
        out = r.dispatch('speech-create', topic='ИИ', minutes=5, themes=('hud',))
        r.dispatch('rollback', entry_id=out['entry'])
        self.assertEqual([p for p in Path(out['folder']).iterdir()], [])

    def test_refusals(self):
        with self.assertRaises(Denied):
            self.router(read_only=True).dispatch('speech-create', topic='ИИ', minutes=5)
        with self.assertRaises(Denied):
            self.router(approve=False).dispatch('speech-create', topic='ИИ', minutes=5)
        with self.assertRaises(Denied):
            self.router().dispatch('speech-create', topic='ИИ', minutes=5, themes=('neon',))
        self.assertFalse((Path(self.tmp.name) / 'presentations').exists())

    def test_schema_is_strict_and_brief_carries_user_facts(self):
        self.assertEqual(set(SLIDE['required']), set(SLIDE['properties']))
        self.assertEqual(set(SCHEMA['required']), set(SCHEMA['properties']))
        self.assertEqual(SLIDE['properties']['kind']['enum'], KINDS)
        b = brief('Тема', 10, audience='логисты', notes='Наш склад: 40 машин')
        self.assertIn('1300–1430 слов', b)
        self.assertIn('40 машин', b)
        with self.assertRaises(Denied):
            validate_plan({'title': 'x', 'subtitle': '', 'speech': [], 'slides': [], 'to_verify': []})


if __name__ == '__main__':
    unittest.main()
