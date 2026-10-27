"""JARVIS local integration block.

Unit: session, containment, router privileges, parsers (no external processes).
Integration: real V8 engine + V8 venv (skipped if V8 is not installed).
Browser: real offline Chromium via existing BrowserSnapshot (skipped if not configured).
Windows E2E: real UIA read-only on the V2 test window, opt-in JARVIS_WINDOWS_E2E=1.
Approvals in automated tests are approval_mode='test-auto', never keyboard.
"""
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import pickle
import tempfile
import time
import unittest
from developer_core.jarvis_local import router as router_mod
from developer_core.jarvis_local.router import DEFAULT_V2, DEFAULT_V8, UnifiedTaskRouter, remote_registration
from developer_core.jarvis_local.session import Denied, TrustedLocalSession
from developer_core.jarvis_local.windows import WindowsOperator, parse_invoke
from developer_core.jarvis_local.workspace import Workspace

V8_READY = ((DEFAULT_V8 / '.venv-ufo-v8' / 'Scripts' / 'python.exe').is_file()
            and importlib.util.find_spec('developer_core.ufo_unified') is not None)  # full JARVIS DEV tree
BROWSER_ENGINE = importlib.util.find_spec('developer_core.browser_snapshot') is not None


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


class Approver:
    def __init__(self, answer=True):
        self.answer, self.seen = answer, []

    def __call__(self, summary, challenge):
        self.seen.append(summary)
        return self.answer


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name) / 'ws'
        (self.root / 'docs').mkdir(parents=True)
        (self.root / 'docs' / 'a.txt').write_text('alpha', encoding='utf-8')
        self.approver = Approver()
        self.router = UnifiedTaskRouter(TrustedLocalSession.for_tests(self.approver), self.root)


class SessionTests(unittest.TestCase):
    def test_redirected_console_refused(self):
        with self.assertRaises(Denied):
            TrustedLocalSession.open_interactive(stdin=io.StringIO(), stdout=io.StringIO(), environ={})

    def test_remote_marker_refused(self):
        class Tty(io.StringIO):
            def isatty(self):
                return True
        with self.assertRaises(Denied):
            TrustedLocalSession.open_interactive(stdin=Tty(), stdout=Tty(), environ={'SSH_CONNECTION': 'x'})
        s = TrustedLocalSession.open_interactive(stdin=Tty(), stdout=Tty(), environ={})
        self.assertEqual(s.approval_mode, 'local-keyboard')

    def test_not_transferable(self):
        with self.assertRaises(Denied):
            pickle.dumps(TrustedLocalSession.for_tests(Approver()))

    def test_read_only_refuses_every_approval(self):
        with self.assertRaises(Denied):
            TrustedLocalSession.read_only().confirm({'task': 'x'})

    def test_declined_and_expired(self):
        with self.assertRaises(Denied):
            TrustedLocalSession.for_tests(Approver(False)).confirm({})
        s = TrustedLocalSession.for_tests(Approver(), ttl=0)
        time.sleep(0.01)
        with self.assertRaises(Denied):
            s.check()


class ContainmentAndRouterTests(Base):
    def test_paths(self):
        ws = self.router.ws
        for bad in ('../x', 'C:/Windows/x', '/etc/x', '.jarvis_local/audit.jsonl', '.ufo_v8_audit.jsonl', 'a/./b', ''):
            with self.assertRaises(Denied, msg=bad):
                ws.resolve(bad)
        self.assertEqual(ws.resolve('docs/a.txt', exists=True), (self.root / 'docs' / 'a.txt').resolve())

    def test_symlink_denied(self):
        target = Path(self.tmp.name) / 'outside'
        target.mkdir()
        try:
            os.symlink(target, self.root / 'link', target_is_directory=True)
        except OSError:
            self.skipTest('symlink creation not permitted')
        with self.assertRaises(Denied):
            self.router.ws.resolve('link/x.txt')

    def test_router_requires_session_and_known_task(self):
        with self.assertRaises(Denied):
            UnifiedTaskRouter(object(), self.root)
        with self.assertRaises(Denied):
            self.router.dispatch('shell', cmd='whoami')
        ro = UnifiedTaskRouter(TrustedLocalSession.read_only(), self.root)
        with self.assertRaises(Denied):
            ro.dispatch('windows-test-window')
        with self.assertRaises(Denied):
            ro.dispatch('browser-allow-host', host='example.com')

    def test_not_registered_on_remote_surfaces(self):
        self.assertEqual(remote_registration(), [])
        fake = Path(self.tmp.name) / 'pkg'
        fake.mkdir()
        (fake / 'workbench.py').write_text('from .jarvis_local import router', encoding='utf-8')
        self.assertEqual(remote_registration(fake), ['workbench.py'])

    def test_every_task_has_handler(self):
        for task in router_mod.TASKS:
            self.assertTrue(hasattr(self.router, '_' + task.replace('-', '_')), task)

    def test_office_plan_rules(self):
        office = self.router.office
        (self.root / 'docs' / 'b.xlsx').write_bytes(b'x')
        with self.assertRaises(Denied):
            office.office_plan('docs/a.txt', 'docs/c.txt', 'docx-append-paragraph', {'text': 'x'})
        with self.assertRaises(Denied):
            office.office_plan('docs/b.xlsx', 'docs/b.xlsx', 'xlsx-set-cell', {'cell': 'A1', 'value': 'v'})
        with self.assertRaises(Denied):
            office.office_plan('docs/b.xlsx', 'docs/c.xlsx', 'xlsx-set-cell', {'cell': 'A1', 'value': '=HYPERLINK("x")'})
        with self.assertRaises(Denied):
            office.office_plan('docs/b.xlsx', 'docs/c.xlsx', 'xlsx-delete-sheet', {})


class WindowsUnitTests(Base):
    def test_parse_transcript(self):
        self.assertEqual(parse_invoke('x\nUFO_V2=PASS before=COUNTER: 2 after=COUNTER: 3\n'),
                         {'outcome': 'PASS', 'before': 2, 'after': 3})
        self.assertEqual(parse_invoke('UFO_V2=CANCELLED')['outcome'], 'CANCELLED')
        self.assertEqual(parse_invoke('')['outcome'], 'NO_RESULT')

    def test_automated_session_can_never_click(self):
        def runner(*a, **k):
            raise AssertionError('no process may start')
        op = WindowsOperator(DEFAULT_V2, Workspace(self.root), runner=runner)
        with self.assertRaises(Denied):
            op.invoke_test_button(TrustedLocalSession.for_tests(Approver()))


class NotepadUnitTests(Base):
    def test_parsers_and_text_rules(self):
        from developer_core.jarvis_local.notepad import check_text, parse_result, text_sha256
        h = 'a' * 64
        self.assertEqual(parse_result('NOTEPAD=PASS pid=5 before_sha256=%s after_sha256=%s length=3' % (h, h))['outcome'], 'PASS')
        self.assertEqual(parse_result('NOTEPAD_ERROR=NOTEPAD=CANCELLED')['outcome'], 'CANCELLED')
        self.assertEqual(parse_result('NOTEPAD_ERROR=DENIED: target changed')['outcome'], 'DENIED')
        self.assertEqual(text_sha256('a\r\nb'), text_sha256('a\nb'))
        for bad in ('', 'x' * 2001, 'a\x07b'):
            with self.assertRaises(Denied):
                check_text(bad)
        self.assertEqual(check_text('Привет\n\tмир'), 'Привет\n\tмир')

    def test_automated_session_can_never_type_and_no_state_no_target(self):
        def runner(*a, **k):
            raise AssertionError('no process may start')
        from developer_core.jarvis_local.notepad import NotepadOperator
        op = NotepadOperator(Workspace(self.root), runner=runner, popen=runner)
        with self.assertRaises(Denied):
            op.set_text('hi', TrustedLocalSession.for_tests(Approver()))
        with self.assertRaises(Denied):
            op.inspect()
        ro = UnifiedTaskRouter(TrustedLocalSession.read_only(), self.root)
        with self.assertRaises(Denied):
            ro.dispatch('notepad-open')


@unittest.skipUnless(V8_READY, 'V8 engine/venv not installed')
class V8IntegrationTests(Base):
    def test_capability_is_real(self):
        status = self.router.dispatch('status')
        self.assertTrue(status['v8_available'], status['v8'])
        self.assertTrue(status['v8']['office_worker_ready'])
        self.assertFalse(status['workbench_registered'])

    def test_search_hides_internal_state(self):
        self.router.dispatch('file-plan', entries=[{'op': 'copy', 'src': 'docs/a.txt', 'dst': 'docs/b.txt'}])
        self.assertEqual(self.router.dispatch('search', pattern='*')['files'], ['docs/a.txt'])

    def test_copy_apply_audit_and_rollback(self):
        entries = [{'op': 'copy', 'src': 'docs/a.txt', 'dst': 'docs/b.txt'}]
        plan = self.router.dispatch('file-plan', entries=entries)
        self.assertEqual(plan['entries'][0]['sha256'], sha(self.root / 'docs' / 'a.txt'))
        result = self.router.dispatch('file-apply', entries=entries, approval=plan['approval'])
        self.assertTrue(result['verified'])
        self.assertEqual((self.root / 'docs' / 'b.txt').read_text(encoding='utf-8'), 'alpha')
        self.assertEqual(self.approver.seen[-1]['task'], 'file-apply')
        entry = self.router.ws.journal.find(result['audit_id'])
        self.assertEqual(entry['approval_mode'], 'test-auto')
        self.router.dispatch('rollback', entry_id=result['audit_id'])
        self.assertFalse((self.root / 'docs' / 'b.txt').exists())
        self.assertTrue((self.root / 'docs' / 'a.txt').exists())
        with self.assertRaises(Denied):
            self.router.dispatch('rollback', entry_id=result['audit_id'])

    def test_move_rollback_restores_source(self):
        entries = [{'op': 'move', 'src': 'docs/a.txt', 'dst': 'docs/moved.txt'}]
        plan = self.router.dispatch('file-plan', entries=entries)
        result = self.router.dispatch('file-apply', entries=entries, approval=plan['approval'])
        self.assertFalse((self.root / 'docs' / 'a.txt').exists())
        self.router.dispatch('rollback', entry_id=result['audit_id'])
        self.assertEqual((self.root / 'docs' / 'a.txt').read_text(encoding='utf-8'), 'alpha')
        self.assertFalse((self.root / 'docs' / 'moved.txt').exists())

    def test_wrong_hash_declined_and_stale_source_write_nothing(self):
        entries = [{'op': 'copy', 'src': 'docs/a.txt', 'dst': 'docs/b.txt'}]
        with self.assertRaises(Denied):
            self.router.dispatch('file-apply', entries=entries, approval='0' * 64)
        plan = self.router.dispatch('file-plan', entries=entries)
        self.approver.answer = False
        with self.assertRaises(Denied):
            self.router.dispatch('file-apply', entries=entries, approval=plan['approval'])
        self.approver.answer = True
        (self.root / 'docs' / 'a.txt').write_text('changed', encoding='utf-8')
        with self.assertRaises(Denied):
            self.router.dispatch('file-apply', entries=entries, approval=plan['approval'])
        self.assertFalse((self.root / 'docs' / 'b.txt').exists())

    def test_existing_target_never_overwritten(self):
        (self.root / 'docs' / 'b.txt').write_text('keep', encoding='utf-8')
        with self.assertRaises(Denied):
            self.router.dispatch('file-plan', entries=[{'op': 'copy', 'src': 'docs/a.txt', 'dst': 'docs/b.txt'}])
        self.assertEqual((self.root / 'docs' / 'b.txt').read_text(encoding='utf-8'), 'keep')

    def test_docx_edit_on_copy(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest('python-docx not in DEV venv for fixtures')
        d = Document()
        d.add_paragraph('Первый абзац')
        d.save(self.root / 'docs' / 'r.docx')
        before = sha(self.root / 'docs' / 'r.docx')
        args = {'src': 'docs/r.docx', 'dst': 'docs/r-copy.docx', 'op': 'docx-append-paragraph',
                'payload': {'text': 'Добавлено JARVIS'}}
        plan = self.router.dispatch('office-plan', **args)
        result = self.router.dispatch('office-apply', **args, approval=plan['approval'])
        self.assertTrue(result['original_unchanged'])
        self.assertEqual(sha(self.root / 'docs' / 'r.docx'), before)
        # Independent check (DEV python-docx) and through the V8 reader: both must keep Cyrillic intact.
        self.assertEqual(Document(self.root / 'docs' / 'r-copy.docx').paragraphs[-1].text, 'Добавлено JARVIS')
        read = self.router.dispatch('office-read', path='docs/r-copy.docx')
        self.assertEqual(read['paragraphs'], ['Первый абзац', 'Добавлено JARVIS'])
        self.router.dispatch('rollback', entry_id=result['audit_id'])
        self.assertFalse((self.root / 'docs' / 'r-copy.docx').exists())

    def test_xlsx_edits_on_copy(self):
        try:
            from openpyxl import Workbook, load_workbook
        except ImportError:
            self.skipTest('openpyxl not in DEV venv for fixtures')
        wb = Workbook()
        wb.active.title = 'Data'
        wb.active.append(['name', 'qty'])
        wb.save(self.root / 'docs' / 't.xlsx')
        before = sha(self.root / 'docs' / 't.xlsx')
        for dst, op, payload in (('docs/t1.xlsx', 'xlsx-append-row', {'cells': ['яблоко', '3']}),
                                 ('docs/t2.xlsx', 'xlsx-set-cell', {'sheet': 'Data', 'cell': 'C1', 'value': 'note'})):
            args = {'src': 'docs/t.xlsx', 'dst': dst, 'op': op, 'payload': payload}
            plan = self.router.dispatch('office-plan', **args)
            self.router.dispatch('office-apply', **args, approval=plan['approval'])
        self.assertEqual(sha(self.root / 'docs' / 't.xlsx'), before)
        self.assertEqual([c.value for c in load_workbook(self.root / 'docs' / 't1.xlsx')['Data'][2]], ['яблоко', '3'])
        rows = self.router.dispatch('office-read', path='docs/t1.xlsx')['sheets'][0]['rows']
        self.assertEqual(rows[-1], ['яблоко', '3'])
        self.assertEqual(load_workbook(self.root / 'docs' / 't2.xlsx')['Data']['C1'].value, 'note')
        leftovers = [p.name for p in (self.root / 'docs').iterdir() if p.name.startswith('.~jarvis_tmp_')]
        self.assertEqual(leftovers, [])


class BrowserTests(Base):
    @unittest.skipUnless(BROWSER_ENGINE, 'browser engine of the full JARVIS DEV tree not installed')
    def test_offline_selftest_or_not_configured(self):
        status = self.router.dispatch('browser-status')
        self.assertFalse(status['interactive_control'])
        if not (status['package_installed'] and status['browser_directory_present'] and status['process_tree_supported']):
            self.skipTest('Playwright/Chromium not configured')
        result = self.router.dispatch('browser-selftest')
        self.assertEqual(result['status'], 'PASS', result)
        self.assertFalse(result['javascript_executed'])

    def test_unlisted_host_denied_without_network(self):
        with self.assertRaises(Denied):
            self.router.dispatch('browser-read', url='https://example.com/')
        with self.assertRaises(Denied):
            self.router.dispatch('browser-read', url='http://example.com/')


@unittest.skipUnless(os.environ.get('JARVIS_WINDOWS_E2E') == '1' and os.name == 'nt', 'set JARVIS_WINDOWS_E2E=1')
class WindowsE2ETests(Base):
    def test_test_window_read_only_uia(self):
        self.assertEqual(self.router.dispatch('windows-check')['status'], 'PASS')
        self.router.dispatch('windows-test-window')
        try:
            found = {'status': 'FAIL'}
            for _ in range(30):
                found = self.router.dispatch('windows-inspect')
                if found['status'] == 'PASS':
                    break
                time.sleep(0.5)
            self.assertEqual(found['status'], 'PASS', found)
            self.assertEqual(found['counter'], 0)
        finally:
            self.assertIn(self.router.dispatch('windows-close-test-window')['status'], ('CLOSED', 'ALREADY_CLOSED'))


if __name__ == '__main__':
    unittest.main()
