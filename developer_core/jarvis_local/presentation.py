"""Speech plan (claude_writer JSON) -> PowerPoint deck + Word speech, built locally.

Three themes drawn with native shapes (no stock images): "hud" (JARVIS dark HUD), "business"
(clean light), "creative" (bold colour blocks). Every slide carries its speech fragment in the
speaker notes. Files are new (never overwrite), live in <workspace>/presentations/<folder>/ and are
journaled with a remove_created rollback.
"""
from __future__ import annotations
import datetime
import json
import re
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt
from .session import Denied
from .workspace import sha256

W, H = Emu(12192000), Emu(6858000)  # 16:9
IMAGE_X = 0.58  # picture slot: right part of the slide, full height
IMAGE_RATIO = (1 - IMAGE_X) * 16 / 9  # width/height of that slot
IMAGE_KINDS = ('title', 'closing', 'section', 'statement', 'bullets')
THEMES = {
    'hud': {'bg': '070B12', 'fg': 'E8F4FF', 'muted': '7FA3BF', 'accent': '29D3FF', 'accent2': 'FFB84D',
            'panel': '0E1724', 'head': 'Bahnschrift SemiBold', 'body': 'Bahnschrift'},
    'business': {'bg': 'FFFFFF', 'fg': '14213D', 'muted': '5C6B80', 'accent': '1F6FEB', 'accent2': 'F2A541',
                 'panel': 'F1F4F9', 'head': 'Segoe UI Semibold', 'body': 'Segoe UI'},
    'creative': {'bg': '14111F', 'fg': 'FFFFFF', 'muted': 'C9C2E0', 'accent': 'FF4F7B', 'accent2': '7C5CFF',
                 'panel': '221C33', 'head': 'Segoe UI Black', 'body': 'Segoe UI'},
}


def _rgb(hexstr):
    return RGBColor.from_string(hexstr)


class Deck:
    def __init__(self, theme='hud', total=0):
        if theme not in THEMES:
            raise Denied('Unknown theme: %s' % theme)
        self.t, self.name = THEMES[theme], theme
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.n, self.total, self.sections = 0, total, 0
        self.right, self.scale = 1.0, 1.0  # text area limit and font scale (narrower next to a picture)

    # ---- primitives ----
    def _slide(self, notes):
        s = self.prs.slides.add_slide(self.blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = _rgb(self.t['bg'])
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        self.n += 1
        return s

    def _box(self, s, shape, x, y, w, h, fill=None, line=None, width=1.0):
        sh = s.shapes.add_shape(shape, Emu(int(x * W)), Emu(int(y * H)), Emu(int(w * W)), Emu(int(h * H)))
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = _rgb(fill)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = _rgb(line)
            sh.line.width = Pt(width)
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    def _text(self, s, text, x, y, w, h, size, color=None, bold=False, head=False, align=PP_ALIGN.LEFT,
              anchor=MSO_ANCHOR.TOP):
        w = min(w, self.right - 0.04 - x)
        size = max(12, int(size * self.scale))
        tb = s.shapes.add_textbox(Emu(int(x * W)), Emu(int(y * H)), Emu(int(w * W)), Emu(int(h * H)))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
            setattr(tf, m, 0)
        lines = text if isinstance(text, list) else [text]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = line
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = self.t['head'] if head else self.t['body']
            r.font.color.rgb = _rgb(color or self.t['fg'])
            if len(lines) > 1:
                p.space_after = Pt(size * 0.55)
        return tb

    @staticmethod
    def _fit(text, big, small, limit):
        return big if len(text) <= limit else max(small, int(big * limit / len(text)))

    def _decor(self, s):
        t = self.t
        if self.name == 'hud':
            lw, lh = 0.0025, 0.0045  # line thickness as slide fractions (about 3 px)
            for left in (True, False):  # corner brackets
                for top in (True, False):
                    x = 0.03 if left else 0.97 - 0.03
                    y = 0.05 if top else 0.95 - lh
                    self._box(s, MSO_SHAPE.RECTANGLE, x, y, 0.03, lh, t['accent'])
                    vx = 0.03 if left else 0.97 - lw
                    vy = 0.05 if top else 0.95 - 0.055
                    self._box(s, MSO_SHAPE.RECTANGLE, vx, vy, lw, 0.055, t['accent'])
            self._text(s, 'J.A.R.V.I.S  //  %02d' % self.n, 0.06, 0.915, 0.4, 0.04, 10, t['muted'])
            ring = self._box(s, MSO_SHAPE.OVAL, 0.80, -0.12, 0.30, 0.53, None, t['accent'], 0.75)
            ring.line.dash_style = 4
        elif self.name == 'business':
            self._box(s, MSO_SHAPE.RECTANGLE, 0, 0, 0.012, 1, t['accent'])
            self._text(s, '%02d' % self.n, 0.92, 0.92, 0.05, 0.04, 11, t['muted'], align=PP_ALIGN.RIGHT)
        else:
            self._box(s, MSO_SHAPE.OVAL, 0.86, -0.30, 0.34, 0.60, t['accent2'])
            self._box(s, MSO_SHAPE.OVAL, 0.93, 0.74, 0.16, 0.28, t['accent'])

    def _progress(self, s):
        """Thin bar at the bottom: where we are in the talk (drawn last, above any picture)."""
        if self.total:
            self._box(s, MSO_SHAPE.RECTANGLE, 0, 0.992, 1, 0.008, self.t['panel'])
            self._box(s, MSO_SHAPE.RECTANGLE, 0, 0.992, self.n / self.total, 0.008, self.t['accent'])

    # ---- slide kinds ----
    def add(self, sl):
        kind = sl['kind']
        notes = sl.get('notes', '') + ('\n\n[%s]' % sl['image_credit'] if sl.get('image_credit') else '')
        s = self._slide(notes)
        self._decor(s)
        t = self.t
        picture = sl.get('image_file') if kind in IMAGE_KINDS else None
        self.right, self.scale = (IMAGE_X, 0.8) if picture else (1.0, 1.0)
        if picture:
            s.shapes.add_picture(picture, Emu(int(IMAGE_X * W)), 0, Emu(int((1 - IMAGE_X) * W)), H)
            self._box(s, MSO_SHAPE.RECTANGLE, IMAGE_X - 0.003, 0, 0.004, 1, t['accent'])
        title = sl.get('title', '')
        if kind in ('title', 'closing'):
            if self.name != 'business':
                self._box(s, MSO_SHAPE.RECTANGLE, 0.08, 0.36, 0.08, 0.008, t['accent'])
            self._text(s, title, 0.08, 0.40, 0.72, 0.30, self._fit(title, 48, 30, 40), head=True, bold=True)
            if sl.get('subtitle'):
                self._text(s, sl['subtitle'], 0.08, 0.72, 0.72, 0.12, 20, t['muted'])
        elif kind == 'section':
            self.sections += 1
            self._text(s, '%02d' % self.sections, 0.08, 0.22, 0.3, 0.2, 72, t['accent'], bold=True, head=True)
            self._text(s, title, 0.08, 0.48, 0.75, 0.25, self._fit(title, 44, 28, 36), head=True, bold=True)
            if sl.get('subtitle'):
                self._text(s, sl['subtitle'], 0.08, 0.74, 0.75, 0.12, 18, t['muted'])
        elif kind == 'statement':
            self._box(s, MSO_SHAPE.RECTANGLE, 0.08, 0.30, 0.006, 0.40, t['accent'])
            self._text(s, title, 0.11, 0.26, 0.74, 0.48, self._fit(title, 40, 26, 70), head=True, bold=True,
                       anchor=MSO_ANCHOR.MIDDLE)
            if sl.get('subtitle'):
                self._text(s, sl['subtitle'], 0.11, 0.78, 0.74, 0.1, 18, t['muted'])
        elif kind == 'number':
            self._text(s, title, 0.08, 0.10, 0.80, 0.12, 26, head=True, bold=True)
            num = sl.get('number') or '—'
            self._text(s, num, 0.08, 0.24, 0.84, 0.42, self._fit(num, 150, 60, 6), t['accent'], bold=True,
                       head=True, anchor=MSO_ANCHOR.MIDDLE)
            self._text(s, sl.get('number_label', ''), 0.08, 0.70, 0.72, 0.16, 22, t['muted'])
        elif kind == 'quote':
            self._text(s, '“', 0.06, 0.08, 0.2, 0.3, 160, t['accent'], bold=True, head=True)
            q = sl.get('quote') or title
            self._text(s, q, 0.12, 0.28, 0.74, 0.42, self._fit(q, 34, 22, 110), head=True)
            if sl.get('author'):
                self._text(s, '— ' + sl['author'], 0.12, 0.74, 0.7, 0.08, 18, t['muted'])
        elif kind == 'compare':
            self._text(s, title, 0.08, 0.08, 0.80, 0.12, 30, head=True, bold=True)
            for i, (col, items) in enumerate(((sl.get('left_title', ''), sl.get('left', [])),
                                              (sl.get('right_title', ''), sl.get('right', [])))):
                x = 0.08 + i * 0.42
                accent = t['muted'] if i == 0 else t['accent']
                self._box(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 0.27, 0.38, 0.60, t['panel'],
                          accent if self.name == 'hud' else None, 1.0)
                self._text(s, col, x + 0.03, 0.31, 0.32, 0.08, 22, accent, bold=True, head=True)
                self._text(s, ['• ' + b for b in items[:4]], x + 0.03, 0.42, 0.32, 0.42, 20)
        else:  # bullets
            self._text(s, title, 0.08, 0.10, 0.70, 0.14, self._fit(title, 34, 24, 40), head=True, bold=True)
            items = sl.get('bullets', [])[:4]
            for i, b in enumerate(items):
                y = 0.32 + i * 0.145
                self._box(s, MSO_SHAPE.OVAL if self.name != 'business' else MSO_SHAPE.RECTANGLE,
                          0.08, y + 0.025, 0.014, 0.025, t['accent'] if i % 2 == 0 else t['accent2'])
                self._text(s, b, 0.11, y, 0.70, 0.12, 22, anchor=MSO_ANCHOR.TOP)
        self.right, self.scale = 1.0, 1.0
        self._progress(s)
        return s

    def save(self, path):
        self.prs.save(str(path))


def build_deck(plan, theme='hud'):
    deck = Deck(theme, total=len(plan['slides']))
    for sl in plan['slides']:
        deck.add(sl)
    return deck


def write_speech_docx(plan, path, minutes):
    from docx import Document
    from docx.shared import Pt as DPt
    doc = Document()
    doc.add_heading(plan['title'], 0)
    if plan.get('subtitle'):
        doc.add_paragraph(plan['subtitle'])
    doc.add_paragraph('Длительность: %d мин · %d слов (~%.1f мин при спокойном темпе)'
                      % (minutes, plan.get('words', 0), plan.get('minutes_estimate', 0)))
    for sec in plan['speech']:
        doc.add_heading(sec['heading'], 1)
        for para in [p for p in sec['text'].split('\n') if p.strip()]:
            doc.add_paragraph(para.strip()).runs[0].font.size = DPt(13)
    if plan.get('to_verify'):
        doc.add_heading('Проверить перед выступлением', 1)
        for item in plan['to_verify']:
            doc.add_paragraph(item, style='List Bullet')
    doc.save(str(path))


def slug(text):
    s = re.sub(r'[^0-9a-zа-яё]+', '-', text.lower()).strip('-')
    return s[:40] or 'speech'


def create(ws, plan, minutes, themes=('hud',), now=None, image_providers=None):
    """Write plan.json, speech.docx, pictures and one deck per theme into a NEW folder; journal + rollback."""
    stamp = (now or datetime.datetime.now()).strftime('%Y%m%d-%H%M')
    folder = ws.resolve('presentations/%s-%s' % (stamp, slug(plan['title'])))
    if folder.exists():
        raise Denied('Folder already exists: ' + ws.rel(folder))
    folder.mkdir(parents=True)
    files, images = [], []
    if image_providers:
        from .slide_images import fetch_all
        images = fetch_all(plan, folder / 'images', image_providers, ratio=IMAGE_RATIO)
        files += sorted((folder / 'images').glob('*.jpg'))
    with (folder / 'plan.json').open('x', encoding='utf-8') as f:
        f.write(json.dumps(plan, ensure_ascii=False, indent=2))
    files.append(folder / 'plan.json')
    tmp = folder / 'speech.docx'
    write_speech_docx(plan, tmp, minutes)
    files.append(tmp)
    for theme in themes:
        p = folder / ('slides-%s.pptx' % theme)
        build_deck(plan, theme).save(p)
        files.append(p)
    rollback = [{'type': 'remove_created', 'path': ws.rel(f), 'sha256': sha256(f)} for f in files]
    entry = ws.journal.record('speech-create', 'PASS', folder=ws.rel(folder), files=[ws.rel(f) for f in files],
                              slides=len(plan['slides']), words=plan.get('words'), rollback=rollback)
    return {'status': 'PASS', 'folder': str(folder), 'files': [ws.rel(f) for f in files], 'entry': entry['id'],
            'words': plan.get('words'), 'minutes_estimate': plan.get('minutes_estimate'),
            'to_verify': plan.get('to_verify', []), 'images': images}
