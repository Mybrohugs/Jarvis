"""Excel operator: the owner's real spreadsheets in the allow-listed folders (Documents, Desktop, Downloads).

Read: a summary of a table, a column total/average/min/max, rows containing a word.
Write (local approval, audited, undoable): create a new table with headers, add a row, open a table in Excel.
Edits keep the original safe: a backup goes to the workspace first, and rollback puts it back only if
nobody changed the file since. openpyxl cannot keep charts, pivot tables, pictures or macros, so a
workbook that has them is never rewritten: the row goes into a new copy "<name> (JARVIS).xlsx".
A table that is open in Excel (its ~$ lock file exists) is not touched.
"""
from __future__ import annotations
import datetime
import os
from pathlib import Path
import re
import shutil
import time
import zipfile
from .explorer import TITLES, Explorer, _safe_name
from .session import Denied
from .workspace import sha256

ROOTS = ('Documents', 'Desktop', 'Downloads')
EXT = ('.xlsx',)
MAX_BYTES = 30 * 1024 ** 2
MAX_ROWS_READ = 100_000
MAX_CELLS = 30
# parts openpyxl drops or damages on save -> write a copy instead of the original
FRAGILE = ('xl/charts/', 'xl/pivotTables/', 'xl/pivotCache/', 'xl/drawings/', 'xl/media/', 'xl/vbaProject',
           'xl/externalLinks/', 'xl/slicers/', 'xl/timelines/', 'xl/model/', 'customXml/')
FUNCS = {'сумма': 'sum', 'среднее': 'avg', 'максимум': 'max', 'минимум': 'min', 'количество': 'count'}


def _number(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, str) and re.fullmatch(r'\s*\d{1,2}\.\d{1,2}(?:\.\d{2,4})?\s*', v):
        return None  # 01.10 / 01.10.2026 is a date written as text, not 1,1
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.replace(' ', '').replace(' ', '').replace(',', '.').rstrip('₽р$€')
        if re.fullmatch(r'-?\d+(?:\.\d+)?', s):
            return float(s)
    return None


def _fmt(x):
    if x is None:
        return '—'
    if abs(x - round(x)) < 1e-9:
        return '{:,}'.format(int(round(x))).replace(',', ' ')
    return ('%.2f' % x).replace('.', ',')


def _cell_value(text):
    """Spoken cell text -> number / date / text. Text starting with '=' stays text (see _put): no formulas."""
    t = str(text).strip()
    if re.fullmatch(r'-?\d[\d\s ]*', t):
        return int(re.sub(r'\s', '', t))
    # 29.10 / 29.10.2026 is a date (Russian decimals use a comma: 3,5)
    m = re.fullmatch(r'(\d{1,2})\.(\d{1,2})(?:\.(\d{2,4}))?', t)
    if m:
        y = int(m.group(3) or datetime.date.today().year)
        y = y + 2000 if y < 100 else y
        try:
            return datetime.date(y, int(m.group(2)), int(m.group(1)))
        except ValueError:
            pass
    if re.fullmatch(r'-?\d[\d\s ]*[.,]\d+', t):
        return float(re.sub(r'\s', '', t).replace(',', '.'))
    return t


def _put(sheet, row, col, value):
    cell = sheet.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value.startswith('='):
        cell.data_type = 's'  # written as text: a phrase never becomes a formula
    if isinstance(value, datetime.date):
        cell.number_format = 'DD.MM.YYYY'
    return cell


class ExcelOps:
    def __init__(self, ws, explorer=None, start=os.startfile, titles=None, sleep=time.sleep):
        self.ws = ws
        self.explorer = explorer or Explorer(ws)
        self.start, self.sleep = start, sleep
        if titles is None:
            from .apps import window_titles as titles
        self.titles = titles

    # ---- finding a table ----
    def find(self, name):
        """A spoken table name -> (root key, path) in Documents, Desktop or Downloads."""
        key, p = self.explorer.find_named(name, EXT, ROOTS, ('таблицу', 'таблиц'))
        if p.stat().st_size > MAX_BYTES:
            raise Denied('Таблица больше 30 МБ, её лучше открыть в Excel')
        return key, p

    @staticmethod
    def _open(path, read_only=True):
        import openpyxl
        try:
            return openpyxl.load_workbook(path, read_only=read_only, data_only=read_only)
        except zipfile.BadZipFile:
            raise Denied('Файл повреждён или это не таблица Excel') from None
        except Exception as exc:  # openpyxl raises many types on odd files
            raise Denied('Не смог прочитать таблицу: %s' % type(exc).__name__) from None

    @staticmethod
    def _sheet(wb, sheet):
        if not sheet:
            return wb[wb.sheetnames[0]]
        for name in wb.sheetnames:
            if name.lower() == str(sheet).lower():
                return wb[name]
        raise Denied('Нет листа «%s». Листы: %s' % (sheet, ', '.join(wb.sheetnames)))

    @staticmethod
    def _rows(ws_):
        rows = []
        for i, row in enumerate(ws_.iter_rows(values_only=True)):
            if i >= MAX_ROWS_READ:
                break
            rows.append(list(row))
        while rows and all(v in (None, '') for v in rows[-1]):
            rows.pop()
        return rows

    @staticmethod
    def _headers(rows):
        return [str(v).strip() if v not in (None, '') else '' for v in (rows[0] if rows else [])]

    # ---- read ----
    def summary(self, name, sheet=''):
        key, p = self.find(name)
        wb = self._open(p)
        try:
            sh = self._sheet(wb, sheet)
            rows = self._rows(sh)
            heads = self._headers(rows)
            data = rows[1:]
            totals = []
            for i, h in enumerate(heads):
                nums = [_number(r[i]) for r in data if i < len(r)]
                nums = [n for n in nums if n is not None]
                if h and nums and len(nums) >= max(1, len(data) // 2):
                    totals.append((h, sum(nums)))
            speech = 'Таблица «%s»: %d %s, столбцы: %s.' % (p.stem, len(data), _plural(len(data), 'строка', 'строки', 'строк'),
                                                            ', '.join(h for h in heads if h)[:200] or 'без заголовков')
            if totals:
                speech += ' Итого: %s.' % ', '.join('%s %s' % (h, _fmt(s)) for h, s in totals[:3])
            if len(wb.sheetnames) > 1:
                speech += ' Листов %d: %s.' % (len(wb.sheetnames), ', '.join(wb.sheetnames[:5]))
            return {'status': 'PASS', 'file': p.name, 'folder': TITLES[key], 'sheet': sh.title, 'rows': len(data),
                    'headers': heads, 'totals': dict(totals), 'speech': speech}
        finally:
            wb.close()

    def column(self, name, column, func='sum', sheet=''):
        if func not in ('sum', 'avg', 'max', 'min', 'count'):
            raise Denied('Умею: сумма, среднее, максимум, минимум, количество')
        key, p = self.find(name)
        wb = self._open(p)
        try:
            rows = self._rows(self._sheet(wb, sheet))
        finally:
            wb.close()
        heads = self._headers(rows)
        want = str(column).strip().lower()
        idx = [i for i, h in enumerate(heads) if h.lower() == want] or \
              [i for i, h in enumerate(heads) if h and want in h.lower()]
        if len(idx) != 1:
            raise Denied(('Нет столбца «%s».' if not idx else 'Несколько столбцов похожи на «%s».') % column +
                         ' Столбцы: %s' % ', '.join(h for h in heads if h)[:200])
        i = idx[0]
        nums = [n for n in (_number(r[i]) for r in rows[1:] if i < len(r)) if n is not None]
        if not nums:
            raise Denied('В столбце «%s» нет чисел' % heads[i])
        value = {'sum': sum(nums), 'avg': sum(nums) / len(nums), 'max': max(nums), 'min': min(nums), 'count': len(nums)}[func]
        word = {'sum': 'Сумма', 'avg': 'Среднее', 'max': 'Максимум', 'min': 'Минимум', 'count': 'Чисел'}[func]
        return {'status': 'PASS', 'file': p.name, 'column': heads[i], 'func': func, 'value': value, 'numbers': len(nums),
                'speech': '%s по столбцу «%s» в таблице «%s»: %s.' % (word, heads[i], p.stem, _fmt(value))}

    def find_rows(self, name, text, sheet=''):
        key, p = self.find(name)
        want = str(text).strip().lower()
        if not 1 <= len(want) <= 100:
            raise Denied('Что искать в таблице?')
        wb = self._open(p)
        try:
            rows = self._rows(self._sheet(wb, sheet))
        finally:
            wb.close()
        heads = self._headers(rows)
        hits = []
        for n, r in enumerate(rows[1:], start=2):
            if any(want in str(v).lower() for v in r if v is not None):
                hits.append({'row': n, 'values': ['' if v is None else str(v) for v in r][:MAX_CELLS]})
        if not hits:
            return {'status': 'PASS', 'rows': [], 'speech': 'В таблице «%s» нет «%s».' % (p.stem, text)}
        first = hits[0]['values']
        pairs = ', '.join('%s %s' % (h, v) if h else v for h, v in zip(heads + [''] * len(first), first) if v)[:200]
        return {'status': 'PASS', 'rows': hits[:50], 'headers': heads,
                'speech': 'Нашёл %d %s с «%s». Первая, строка %d: %s.' % (
                    len(hits), _plural(len(hits), 'строку', 'строки', 'строк'), text, hits[0]['row'], pairs)}

    # ---- write ----
    def _locked(self, p):
        return (p.parent / ('~$' + p.name)).exists()

    @staticmethod
    def _fragile(p):
        with zipfile.ZipFile(p) as z:
            return sorted({part for part in FRAGILE for n in z.namelist() if n.startswith(part)})

    def create(self, name, headers, session, folder='Documents'):
        if folder not in ROOTS:
            raise Denied('Создаю таблицы только в Документах, на Рабочем столе или в Загрузках')
        title = _safe_name(re.sub(r'\.xlsx$', '', str(name).strip(' «»"'), flags=re.I))
        heads = [str(h).strip()[:100] for h in headers if str(h).strip()][:MAX_CELLS]
        if not heads:
            raise Denied('Назовите столбцы: «создай таблицу Расходы с колонками Дата, Сумма, Категория»')
        base = self.explorer.root(folder)
        dst = base / (title + '.xlsx')
        if dst.exists():
            raise Denied('Таблица «%s» уже есть в «%s»' % (dst.name, TITLES[folder]))
        approved = session.confirm({'task': 'excel-create', 'file': dst.name, 'folder': TITLES[folder], 'columns': heads})
        import openpyxl
        from openpyxl.styles import Font
        wb = openpyxl.Workbook()
        sh = wb.active
        sh.title = 'Лист1'
        for col, h in enumerate(heads, start=1):
            _put(sh, 1, col, h)
        for c in sh[1]:
            c.font = Font(bold=True)
        sh.freeze_panes = 'A2'
        for i, h in enumerate(heads, start=1):
            sh.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(12, min(40, len(h) + 4))
        tmp = base / ('~jarvis-%d.tmp.xlsx' % os.getpid())
        wb.save(tmp)
        try:
            os.rename(tmp, dst)  # fails instead of overwriting
        except OSError:
            tmp.unlink()
            raise Denied('Таблица «%s» появилась, пока я её создавал' % dst.name) from None
        ok = self._headers(self._read_back(dst)) == heads
        entry = self.ws.journal.record('excel-create', 'PASS' if ok else 'VERIFY_FAILED', root=folder, file=dst.name,
                                       columns=len(heads), rollback=[{'type': 'explorer_remove_copy', 'root': folder,
                                                                      'path': dst.name, 'sha256': sha256(dst)}], **approved)
        return {'status': entry['status'], 'audit_id': entry['id'], 'path': str(dst), 'file': dst.name,
                'speech': 'Создал таблицу «%s» в «%s», столбцы: %s.' % (title, TITLES[folder], ', '.join(heads))}

    def _read_back(self, p, sheet=''):
        wb = self._open(p)
        try:
            return self._rows(self._sheet(wb, sheet))
        finally:
            wb.close()

    def append(self, name, cells, session, sheet=''):
        key, p = self.find(name)
        values = [_cell_value(c) for c in cells if str(c).strip() != ''][:MAX_CELLS]
        if not values:
            raise Denied('Что добавить? Например: добавь строку «29.10; 500; еда» в таблицу Расходы')
        if self._locked(p):
            raise Denied('Таблица «%s» открыта в Excel: закройте её, и я добавлю строку' % p.stem)
        fragile = self._fragile(p)
        before = self._read_back(p, sheet)
        heads = self._headers(before)
        if heads and len(values) > len([h for h in heads if h]) and any(heads):
            raise Denied('В строке %d значений, а столбцов %d (%s)' % (len(values), len([h for h in heads if h]),
                                                                       ', '.join(h for h in heads if h)))
        target = p if not fragile else self.explorer._free_name(p.parent, p.stem + ' (JARVIS).xlsx')
        approved = session.confirm({'task': 'excel-append', 'file': p.name, 'folder': TITLES[key], 'row': [str(v) for v in values],
                                    'writes': 'original (backup kept)' if not fragile else 'new copy ' + target.name,
                                    'why_copy': fragile})
        backup = None
        if not fragile:
            backup = self.ws.state / 'excel_backups' / ('%d-%s' % (int(time.time() * 1000), p.name))
            backup.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(p, backup)
            if sha256(backup) != sha256(p):
                raise Denied('Резервная копия не совпала, ничего не менял')
        wb = self._open(p, read_only=False)
        try:
            sh = self._sheet(wb, sheet)
            last = len(before)  # first empty row after the data (openpyxl max_row counts formatted blanks)
            for col, v in enumerate(values, start=1):
                _put(sh, last + 1, col, v)
            tmp = p.parent / ('~jarvis-%d.tmp.xlsx' % os.getpid())
            wb.save(tmp)
        finally:
            wb.close()
        after = self._read_back(tmp, sheet)
        if len(after) != len(before) + 1 or after[:len(before)] != before:
            tmp.unlink()
            raise Denied('Проверка не прошла: старые строки изменились бы. Ничего не менял')
        if target == p:
            os.replace(tmp, p)
            rollback = [{'type': 'excel_restore', 'root': key, 'path': p.relative_to(self.explorer.root(key)).as_posix(),
                         'backup': str(backup), 'sha256': sha256(p)}]
        else:
            os.rename(tmp, target)
            rollback = [{'type': 'explorer_remove_copy', 'root': key,
                         'path': str(target.relative_to(self.explorer.root(key))).replace('\\', '/'),
                         'sha256': sha256(target)}]
        entry = self.ws.journal.record('excel-append', 'PASS', root=key, file=target.name, row=len(after),
                                       copy=bool(fragile), rollback=rollback, **approved)
        speech = 'Добавил строку %d в таблицу «%s».' % (len(after), p.stem)
        if fragile:
            speech = ('В таблице «%s» есть диаграммы или макросы, поэтому оригинал не трогал: строку добавил в копию «%s».'
                      % (p.stem, target.name))
        return {'status': entry['status'], 'audit_id': entry['id'], 'path': str(target), 'row': len(after),
                'copy': bool(fragile), 'speech': speech + ' Отменить: «отмени последнее действие».'}

    def open(self, name):
        key, p = self.find(name)
        self.start(str(p))
        ok = None
        for _ in range(16):
            if any(p.stem.lower()[:30] in t.lower() for t in self.titles()):
                ok = True
                break
            self.sleep(0.5)
        self.ws.journal.record('excel-open', 'PASS' if ok else 'WARN', file=p.name, verified=bool(ok), rollback=[])
        return {'status': 'PASS' if ok else 'WARN', 'verified': bool(ok), 'path': str(p),
                'speech': ('Открыл таблицу «%s».' if ok else 'Открываю таблицу «%s», окно Excel пока не подтвердилось.') % p.stem}

    # ---- rollback (called through Explorer's rollback for real folders) ----
    def plan_restore(self, step):
        p = self.explorer.path(step['root'], step['path'], must_exist='file')
        backup = Path(step['backup'])
        if not backup.is_file() or not backup.resolve().is_relative_to(self.ws.state.resolve()):
            raise Denied('Резервная копия пропала, верните вручную')
        if sha256(p) != step['sha256']:
            raise Denied('Таблицу уже меняли после JARVIS, не откатываю: %s' % p.name)
        if self._locked(p):
            raise Denied('Таблица открыта в Excel: закройте её, и я откачу')
        return p, backup

    @staticmethod
    def restore(p, backup):
        tmp = p.with_name('~jarvis-restore.tmp')
        shutil.copy2(backup, tmp)
        if sha256(tmp) != sha256(backup):
            tmp.unlink()
            raise Denied('Rollback copy verification failed')
        os.replace(tmp, p)
        return {'restored': p.name}


def _plural(n, one, few, many):
    if n % 10 == 1 and n % 100 != 11:
        return one
    if 2 <= n % 10 <= 4 and not 12 <= n % 100 <= 14:
        return few
    return many
