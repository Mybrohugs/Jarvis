"""One-shot smoke for the JARVIS local integration block (no keyboard needed).

Runs in a fresh temp workspace with approval_mode='test-auto' (labelled as such; not a
keyboard approval). Real: V8 engine + V8 venv, Office copy edits, rollback, offline
Chromium, V2 syntax check. Optional --windows-e2e: real read-only UIA on the V2 test
window (opens and closes it). The real button click is NEVER performed here.
Prints JARVIS_LOCAL_SMOKE=PASS|FAIL; exit code 0/1.
"""
import argparse
import json
from pathlib import Path
import sys
import tempfile
import time

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from developer_core.jarvis_local.router import UnifiedTaskRouter  # noqa: E402
from developer_core.jarvis_local.session import TrustedLocalSession  # noqa: E402


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--windows-e2e', action='store_true')
    a = ap.parse_args()
    results = []

    def step(name, kind, fn):
        try:
            detail = fn()
            ok = bool(detail.pop('_ok', True))
        except Exception as exc:  # report every failure mode, keep going
            ok, detail = False, {'error': type(exc).__name__ + ': ' + str(exc)[:300]}
        results.append({'step': name, 'kind': kind, 'result': 'PASS' if ok else 'FAIL', **detail})

    with tempfile.TemporaryDirectory(prefix='jarvis-smoke-') as tmp:
        root = Path(tmp)
        (root / 'docs').mkdir()
        (root / 'docs' / 'note.txt').write_text('smoke', encoding='utf-8')
        from docx import Document
        from openpyxl import Workbook
        d = Document()
        d.add_paragraph('Smoke source')
        d.save(root / 'docs' / 'r.docx')
        wb = Workbook()
        wb.active.append(['a', 'b'])
        wb.save(root / 'docs' / 't.xlsx')
        r = UnifiedTaskRouter(TrustedLocalSession.for_tests(lambda s, c: True), root)

        def status():
            s = r.dispatch('status')
            return {'_ok': s['v8_available'] and not s['workbench_registered'], 'v8_available': s['v8_available'],
                    'packages': s['v8']['packages'], 'workbench_registered': s['workbench_registered']}

        def file_copy():
            e = [{'op': 'copy', 'src': 'docs/note.txt', 'dst': 'docs/note-copy.txt'}]
            res = r.dispatch('file-apply', entries=e, approval=r.dispatch('file-plan', entries=e)['approval'])
            rb = r.dispatch('rollback', entry_id=res['audit_id'])
            return {'_ok': res['verified'] and not (root / 'docs' / 'note-copy.txt').exists(), 'rollback': rb['status']}

        def office(src, dst, op, payload):
            def run():
                args = {'src': src, 'dst': dst, 'op': op, 'payload': payload}
                res = r.dispatch('office-apply', **args, approval=r.dispatch('office-plan', **args)['approval'])
                return {'_ok': res['original_unchanged'] and res['verified_by_reload'], 'created': res['created']}
            return run

        def windows_e2e():
            r.dispatch('windows-test-window')
            try:
                found = {'status': 'FAIL'}
                for _ in range(30):
                    found = r.dispatch('windows-inspect')
                    if found['status'] == 'PASS':
                        break
                    time.sleep(0.5)
                return {'_ok': found['status'] == 'PASS', 'counter': found.get('counter')}
            finally:
                r.dispatch('windows-close-test-window')

        step('status/V8 capability', 'integration(real V8 venv)', status)
        step('file copy + verify + rollback', 'integration(real V8, test-auto approval)', file_copy)
        step('docx append on copy', 'integration(real V8 venv worker)',
             office('docs/r.docx', 'docs/r2.docx', 'docx-append-paragraph', {'text': 'JARVIS проверка'}))
        step('xlsx append row on copy', 'integration(real V8 venv worker)',
             office('docs/t.xlsx', 'docs/t2.xlsx', 'xlsx-append-row', {'cells': ['x', 'y']}))
        step('browser offline Chromium', 'e2e(real Chromium, no network)',
             lambda: (lambda x: {'_ok': x['status'] == 'PASS', **x})(r.dispatch('browser-selftest')))
        step('windows V2 syntax', 'static(PowerShell parser)',
             lambda: (lambda x: {'_ok': x['status'] == 'PASS'})(r.dispatch('windows-check')))
        if a.windows_e2e:
            step('windows test-window UIA read-only', 'e2e(real UIA, test window only)', windows_e2e)
        step('audit journal', 'integration', lambda: {'entries': len(r.dispatch('audit', limit=100)['entries'])})

    ok = all(x['result'] == 'PASS' for x in results)
    print(json.dumps({'steps': results, 'not_run': ['real UI click (needs local keyboard: windows-invoke-test)',
                                                    'browser-read of a live site (needs allow-host approval + network)']},
                     ensure_ascii=False, indent=2))
    print('JARVIS_LOCAL_SMOKE=' + ('PASS' if ok else 'FAIL'))
    return 0 if ok else 1


if __name__ == '__main__':
    raise SystemExit(main())
